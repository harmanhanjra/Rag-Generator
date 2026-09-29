from pathlib import Path

from docx import Document


def create_pdf(path: Path, lines: list[str]) -> None:
    stream = "BT /F1 12 Tf 50 740 Td " + " 0 -22 Td ".join(
        f"({line.replace('(', '[').replace(')', ']')}) Tj" for line in lines
    ) + " ET"
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        f"<< /Length {len(stream)} >>\nstream\n{stream}\nendstream".encode(),
    ]
    output = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for index, content in enumerate(objects, 1):
        offsets.append(len(output))
        output.extend(f"{index} 0 obj\n".encode() + content + b"\nendobj\n")
    xref = len(output)
    output.extend(f"xref\n0 {len(offsets)}\n0000000000 65535 f \n".encode())
    for offset in offsets[1:]:
        output.extend(f"{offset:010d} 00000 n \n".encode())
    output.extend(f"trailer\n<< /Size {len(offsets)} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF".encode())
    path.write_bytes(output)


def create_fixtures(folder: Path) -> None:
    folder.mkdir(parents=True, exist_ok=True)
    create_pdf(folder / "orion-plan.pdf", [
        "Orion project launch plan",
        "The Orion launch date is 14 November 2026.",
        "The Orion project budget is USD 42000.",
        "The Orion project owner is Avery Stone.",
    ])
    (folder / "orion-support.txt").write_text(
        "Orion support policy: The warranty lasts 18 months. Support is available Monday through Friday.", encoding="utf-8")
    (folder / "orion-checklist.md").write_text(
        "# Orion release checklist\nThe approval code is ORION-731. The backup region is North Harbor.", encoding="utf-8")
    doc = Document()
    doc.add_heading("Orion shipping policy", 0)
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Delivery window"
    table.cell(0, 1).text = "Orion delivery takes 5 business days."
    doc.save(folder / "orion-shipping.docx")
    create_pdf(folder / "cedar-plan.pdf", [
        "Cedar project launch plan",
        "The Cedar launch date is 22 December 2026.",
        "The Cedar project budget is USD 97000.",
        "The Cedar project owner is Jordan Vale.",
    ])
    (folder / "cedar-support.txt").write_text(
        "Cedar support policy: The warranty lasts 36 months. The approval code is CEDAR-928.", encoding="utf-8")
    (folder / "orion-addendum.txt").write_text(
        "Orion emergency contact extension is 6732. The escalation channel is Beacon.", encoding="utf-8")
    (folder / "empty.txt").write_text("", encoding="utf-8")
    (folder / "broken.pdf").write_bytes(b"This is not a PDF")
    (folder / "unsupported.csv").write_text("name,value\nexample,42\n", encoding="utf-8")


if __name__ == "__main__":
    create_fixtures(Path(__file__).resolve().parent / "fixtures")
    print("Created synthetic document fixtures")
